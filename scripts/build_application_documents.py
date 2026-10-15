"""지원 문서 생성. 원문과 근거는 assets/documents/application-content.json에서 관리한다."""
import json
import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets/documents"


def base_document(title, name):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.1)
    # Word 기본 템플릿의 파란 제목 테두리를 제거해 지원 문서의 흑백 서식을 유지한다.
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for style_name, size in [("Normal", 11), ("Title", 24), ("Heading 1", 15)]:
        style = document.styles[style_name]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.paragraph_format.space_after = Pt(12)
    document.styles["Normal"].paragraph_format.line_spacing = 1.35
    document.core_properties.author = name
    document.core_properties.title = title
    document.add_paragraph(title, "Title")
    return document


def build(only=None):
    data = json.loads((OUTPUT / "application-content.json").read_text(encoding="utf-8"))
    introduction = base_document(f"{data['name']} 자기소개서", data["name"])
    # Word의 한글 문서 격자는 배수 줄 간격을 격자 단위로 올림해 큰 공백을 만든다.
    # 자기소개서에만 격자 맞춤을 해제하고 포인트 단위 줄 간격을 지정한다.
    for grid in introduction.sections[0]._sectPr.findall(qn("w:docGrid")):
        introduction.sections[0]._sectPr.remove(grid)
    for style_name, size, line in [("Normal", 11, 16), ("Title", 22, 29), ("Heading 1", 13, 19)]:
        style = introduction.styles[style_name]
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = Pt(line)
        style.paragraph_format.keep_with_next = style_name != "Normal"
        style.paragraph_format.keep_together = False
        style.paragraph_format.widow_control = True
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        style.element.get_or_add_pPr().append(snap)
    if data["introduction"].get("field"):
        introduction.add_paragraph(f"지원 분야  {data['introduction']['field']}")
    if data["introduction"]["title"]:
        introduction.add_paragraph(data["introduction"]["title"], "Heading 1")
    for paragraph in data["introduction"]["paragraphs"]:
        introduction.add_paragraph(paragraph)
    for index, scene in enumerate(data["introduction"].get("scenes", [])):
        heading = introduction.add_paragraph(scene["title"], "Heading 1")
        # 제목과 첫 문단만 연결하고 장면 전체를 묶지 않아 큰 빈 공간을 피한다.
        heading.paragraph_format.page_break_before = False
        heading.paragraph_format.space_before = Pt(14 if index else 4)
        heading.paragraph_format.space_after = Pt(5)
        date = introduction.add_paragraph(scene["dateLabel"])
        date.paragraph_format.keep_with_next = True
        date.paragraph_format.space_after = Pt(9)
        for run in date.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("666666")
        paragraphs = scene["body"].split("\n\n")
        for text in paragraphs:
            paragraph = introduction.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.paragraph_format.keep_with_next = False
        ability = introduction.add_paragraph()
        ability.paragraph_format.space_after = Pt(4)
        ability.add_run(scene["ability"]).bold = True
        ability.add_run(" · " + scene["abilitySummary"])
        for run in ability.runs:
            run.font.size = Pt(9)
    introduction.save(OUTPUT / "self-introduction.docx")
    if only == "introduction":
        print("Created self-introduction.docx")
        return

    experience = base_document(f"{data['name']} 경력기술서", data["name"])
    for key in ["scope", "role"]:
        if data["experience"][key]:
            experience.add_paragraph(data["experience"][key])
    # 내용을 나중에 채우는 동안에는 표와 기준일도 없는 빈 문서를 제공한다.
    if not data["experience"]["projects"]:
        experience.save(OUTPUT / "project-experience.docx")
        print("Created self-introduction.docx and project-experience.docx")
        return
    table = experience.add_table(rows=1, cols=5)
    table.autofit = False
    table.style = "Table Grid"
    widths = [2.8, 2.35, 3.3, 4.65, 3.7]
    headings = ["과제와 제출일", "능력", "상황", "행동", "결과"]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, width, heading in zip(table.rows[0].cells, widths, headings):
        cell.width = Cm(width)
        cell.text = heading
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), "F2F2F2")
        cell._tc.get_or_add_tcPr().append(fill)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for project in data["experience"]["projects"]:
        cells = table.add_row().cells
        values = [f"{project['code']}\n{project['title']}\n{project['submitted']}",
                  project["ability"], project["situation"], project["action"], project["result"]]
        for cell, width, value in zip(cells, widths, values):
            cell.width = Cm(width)
            cell.text = value
        # 과제별 상황·행동·결과는 같은 행에 남겨 근거의 연결을 유지한다.
        keep_row = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(keep_row)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
    experience.add_paragraph(f"자료 기준일  {data['asOf']}")
    experience.save(OUTPUT / "project-experience.docx")
    print("Created self-introduction.docx and project-experience.docx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=["introduction"], help="자기소개서만 생성하고 경력기술서는 유지")
    build(parser.parse_args().only)
